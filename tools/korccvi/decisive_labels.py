#!/usr/bin/env python3
"""
2단계 평가 라벨(첫 결정적 지시) 모으기·판정 양식·확정 (docs/labeling/08_label_decisive.md).

  py tools/korccvi/decisive_labels.py agree --labels <라벨러 CSV 폴더> --pack <make_label_pack.py --out 폴더> \
      --calls <import_calls.py --out>/test --out <결과 폴더>
  py tools/korccvi/decisive_labels.py finalize --pack ... --work <agree의 --out> --adjudication <채운 xlsx> \
      --out data/eval/korccvi_labels

라벨러 CSV: 이름에 part1..4 와 chatgpt|gemini|grok 이 들어간 파일 (.csv 또는 .md). call,decisive_seq,kind,reason
  - decisive_seq가 비면 none (kind 칸 값과 상관없이)
  - 형식 밖 줄(메모, "판단하기 어려웠던 점")은 무시
규칙: 세 모델 중 둘 이상 같은 값(번호 또는 none)이면 그 값, 셋이 모두 다르면 사용자 판정.
저장소에 넣는 것: 모델별 (call_id, 번호, kind)와 최종 라벨 — reason은 원문을 인용할 수 있어 넣지 않음.
"""
import argparse
import csv
import io
import json
import os
import re
import sys
from collections import Counter

MODELS = ("chatgpt", "gemini", "grok")
KINDS = {"이체", "현금", "대출", "앱", "정보", "none"}


def read_labels(folder):
    labels = {m: {} for m in MODELS}
    seen = set()
    for f in sorted(os.listdir(folder)):
        mt = re.search(r"part(\d).*?(chatgpt|gemini|grok)", f, re.I)
        if not mt:
            continue
        model, part = mt.group(2).lower(), mt.group(1)
        if (model, part) in seen:
            sys.exit(f"[오류] 같은 모델·part 파일이 둘: {model} part{part}")
        seen.add((model, part))
        with open(os.path.join(folder, f), encoding="utf-8-sig") as fh:
            for r in csv.reader(io.StringIO(fh.read())):
                if not r or not re.fullmatch(r"T\d{3}", r[0].strip()):
                    continue
                t = r[0].strip()
                seq = (r[1] if len(r) > 1 else "").strip()
                kind = (r[2] if len(r) > 2 else "").strip()
                reason = (r[3] if len(r) > 3 else "").strip()
                if t in labels[model]:
                    sys.exit(f"[오류] {model} {t} 두 번 나옴")
                if seq and not seq.isdigit():
                    sys.exit(f"[오류] {model} {t} 번호 이상: {seq}")
                labels[model][t] = {"seq": int(seq) if seq else "none", "kind": kind if seq else "none",
                                    "reason": reason}
    return labels


def load_pack(pack, calls_dir):
    with open(os.path.join(pack, "blind_key.csv"), encoding="utf-8") as f:
        key = {r["call"]: r["call_id"] for r in csv.DictReader(f)}
    calls = {}
    for t, cid in key.items():
        with open(os.path.join(calls_dir, cid + ".jsonl"), encoding="utf-8") as f:
            rows = [json.loads(l) for l in f if l.strip()][1:]
        calls[t] = {r["seq"]: r for r in rows}
    return key, calls


def decide(labels, key):
    out = {}
    for t in sorted(key):
        vals = [labels[m][t]["seq"] for m in MODELS]
        top, n = Counter(vals).most_common(1)[0]
        out[t] = (top, "다수결") if n >= 2 else (None, "판정")
    return out


def agree(a):
    labels = read_labels(a.labels)
    key, calls = load_pack(a.pack, a.calls)
    problems = []
    for m in MODELS:
        miss = set(key) - set(labels[m])
        if miss:
            problems.append(f"{m}: 빠진 통화 {sorted(miss)}")
        for t, v in labels[m].items():
            if v["kind"] not in KINDS:
                problems.append(f"{m} {t}: kind 값 이상 '{v['kind']}'")
            if v["seq"] != "none":
                row = calls[t].get(v["seq"])
                if row is None:
                    problems.append(f"{m} {t}: 통화에 없는 번호 {v['seq']}")
                elif row.get("is_user"):
                    problems.append(f"{m} {t}: \"나\" 줄 번호 {v['seq']} (화자 뒤바뀜 표시 여부: "
                                    f"{'있음' if '뒤바뀜' in v['reason'] else '없음'})")
    if any("빠진 통화" in p for p in problems):
        sys.exit("\n".join(problems))
    dec = decide(labels, key)
    os.makedirs(a.out, exist_ok=True)
    c = Counter(Counter(labels[m][t]["seq"] for m in MODELS).most_common(1)[0][1] for t in key)
    n_none = sum(1 for v, _ in dec.values() if v == "none")
    lines = ["# 첫 결정적 지시 라벨 일치", "", f"- 통화 {len(key)}건, 모델 {', '.join(MODELS)}",
             f"- 세 모델 일치 {c[3]}건, 둘 일치 {c[2]}건, 셋 다름 {c[1]}건 (사용자 판정)",
             f"- 다수결 결과: none {n_none}건, 번호 {sum(1 for v, _ in dec.values() if isinstance(v, int))}건",
             f"- none 여부(지시가 있나 없나)가 모델끼리 갈린 통화: "
             f"{sum(1 for t in key if len({labels[m][t]['seq'] == 'none' for m in MODELS}) > 1)}건", "",
             "## 모델별", "", "| 모델 | none | 번호 |", "|---|---|---|"]
    for m in MODELS:
        nn = sum(1 for v in labels[m].values() if v["seq"] == "none")
        lines.append(f"| {m} | {nn} | {len(labels[m]) - nn} |")
    lines += ["", "## 둘씩 같은 값 비율", "", "| 쌍 | 같음 |", "|---|---|"]
    for i in range(3):
        for k in range(i + 1, 3):
            same = sum(1 for t in key if labels[MODELS[i]][t]["seq"] == labels[MODELS[k]][t]["seq"])
            lines.append(f"| {MODELS[i]}–{MODELS[k]} | {same}/{len(key)} |")
    lines += ["", "## 점검", ""] + [f"- {p}" for p in problems or ["문제 없음"]]
    with open(os.path.join(a.out, "agreement_report.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    with open(os.path.join(a.out, "model_labels.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["call_id"] + [f"{m}_{x}" for m in MODELS for x in ("seq", "kind")])
        for t in sorted(key):
            w.writerow([key[t]] + [labels[m][t][x] for m in MODELS for x in ("seq", "kind")])
    with open(os.path.join(a.out, "labels.json"), "w", encoding="utf-8") as f:
        json.dump(labels, f, ensure_ascii=False)
    make_xlsx(labels, key, calls, dec, os.path.join(a.out, "adjudication_decisive.xlsx"))
    print("\n".join(lines[:6]))
    print(f"점검 문제 {len(problems)}개 → {a.out}")


def make_xlsx(labels, key, calls, dec, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "판정"
    heads = ["통화", "후보", "모델", "kind", "이유", "그 메시지와 앞 두 줄", "판정", "메모"]
    for j, (h, wd) in enumerate(zip(heads, [7, 8, 9, 7, 30, 70, 10, 24]), 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True)
        ws.column_dimensions[c.column_letter].width = wd
    r = 2
    for t in sorted(key):
        if dec[t][0] is not None:
            continue
        cands = {}
        for m in MODELS:
            cands.setdefault(labels[m][t]["seq"], []).append(m)
        opts = [str(s) for s in cands] + ["none"]
        dv = DataValidation(type="list", formula1='"' + ",".join(dict.fromkeys(opts)) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        first = r
        for s, ms in cands.items():
            ctx = ""
            if s != "none":
                seqs = sorted(calls[t])
                i = seqs.index(s) if s in seqs else None
                if i is not None:
                    ctx = "\n".join(f"{q} {'나' if calls[t][q].get('is_user') else '상대'}: {calls[t][q]['text']}"
                                    for q in seqs[max(0, i - 2):i + 1])
                else:
                    ctx = "(통화에 없는 번호)"
            vals = [t, s, ", ".join(ms), labels[ms[0]][t]["kind"], labels[ms[0]][t]["reason"], ctx]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=j, value=v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        cell = ws.cell(row=first, column=7)
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
        dv.add(cell)
        r += 1
    g = wb.create_sheet("안내")
    for i, t in enumerate([
            "셋이 모두 다른 통화만 모았습니다. 통화마다 첫 줄의 '판정' 칸에 번호 또는 none을 고릅니다.",
            "기준: 상대가 피해로 바로 이어지는 행동(이체·현금·대출·앱·정보)을 처음 구체적으로 시킨 메시지 (08_label_decisive.md).",
            "후보 문맥만으로 정하기 어려우면 가림판(blind_partN.md)에서 그 통화 전체를 봅니다.",
            "후보 중 어느 것도 맞지 않으면 메모에 맞는 번호를 적고 판정 칸은 비워 둡니다."], 1):
        g.cell(row=i, column=1, value=t)
    g.column_dimensions["A"].width = 120
    wb.save(path)


def finalize(a):
    from openpyxl import load_workbook
    with open(os.path.join(a.work, "labels.json"), encoding="utf-8") as f:
        labels = json.load(f)
    with open(os.path.join(a.pack, "blind_key.csv"), encoding="utf-8") as f:
        key = {r["call"]: r["call_id"] for r in csv.DictReader(f)}
    dec = decide(labels, key)
    ws = load_workbook(a.adjudication, data_only=True)["판정"]
    adj = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        t, choice, memo = row[0], row[6], row[7]
        if t and (choice not in (None, "") or memo):
            v = str(choice).strip() if choice not in (None, "") else None
            if v is None:
                mm = re.search(r"\d+", str(memo))
                if not mm:
                    sys.exit(f"[오류] {t}: 판정도 메모 번호도 없음")
                v = mm.group(0)
            adj[t] = "none" if v == "none" else int(float(v))
    rows = []
    for t in sorted(key):
        v, src = dec[t]
        if v is None:
            if t not in adj:
                sys.exit(f"[오류] {t}: 판정이 비어 있음")
            v, src = adj[t], "판정"
        kinds = [labels[m][t]["kind"] for m in MODELS if labels[m][t]["seq"] == v]
        kind = Counter(kinds).most_common(1)[0][0] if kinds else ""
        rows.append({"call_id": key[t], "decisive_seq": v, "kind": kind, "source": src})
    os.makedirs(a.out, exist_ok=True)
    with open(os.path.join(a.out, "final_decisive.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["call_id", "decisive_seq", "kind", "source"], lineterminator="\n")
        w.writeheader()
        w.writerows(sorted(rows, key=lambda r: r["call_id"]))
    print(f"최종 {len(rows)}건 (none {sum(1 for r in rows if r['decisive_seq'] == 'none')}, 판정 {len(adj)}) → {a.out}")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    g = sub.add_parser("agree")
    g.add_argument("--labels", required=True)
    g.add_argument("--pack", required=True)
    g.add_argument("--calls", required=True)
    g.add_argument("--out", required=True)
    f = sub.add_parser("finalize")
    f.add_argument("--pack", required=True)
    f.add_argument("--work", required=True)
    f.add_argument("--adjudication", required=True)
    f.add_argument("--out", required=True)
    a = ap.parse_args()
    agree(a) if a.cmd == "agree" else finalize(a)


if __name__ == "__main__":
    main()
