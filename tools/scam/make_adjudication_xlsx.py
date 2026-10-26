#!/usr/bin/env python3
"""
최종 라벨 판정용 엑셀. 모델 다수결로 정하지 않는 칸(갈린 칸, 기준표 v2.1에서 규칙이 바뀐 칸)만 모아
세 모델의 답과 해당 기준표 문장을 나란히 보여 주고, 사람이 고르게 한다.

사용법
  python tools/scam/make_adjudication_xlsx.py data/eval/scam/*.jsonl --shuffle-seed 20261021 \
      --items docs/labeling/label_results/v2/adjudication_items.csv \
      --labels docs/labeling/label_results/v2 --out <파일.xlsx>
  홀드아웃(평가 세트 v2): --models chatgpt,gemini,grok --guide "기준표 v2.2" \
      --blind-name blind_cases.md (기본값은 기존 평가 세트 2차 라벨링 값)

items CSV: case, seq, field, rule
  field = stage | sig:<코드> | point:<first_suggest|first_intent_before|first_transfer_before|transfer_seen>
  rule  = 06_label_round2.md 3장 번호 (RULES에 기준표 v2.1 문장)
labels 폴더: labels_<모델>_messages.csv / labels_<모델>_points.csv (--models, 기본 grok, claude, chatgpt)
"""
import argparse
import csv
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from agreement import read_csv, parse_signals  # noqa: E402
from make_blind import display_sender, room_names  # noqa: E402
from make_label_sheet import ordered  # noqa: E402
from make_label_xlsx import (BASE, BOLD, CENTER, HEAD_FILL, INPUT_FILL, SIGNALS, STAGES, TITLE,  # noqa: E402
                             WRAP, style_cell)

MODELS = ("grok", "claude", "chatgpt")
SIGNAL_NAME = {c: n for c, n, _ in SIGNALS}
POINT_NAME = {"first_suggest": "첫 금전 권유", "first_intent_before": "첫 송금 의사 직전",
              "first_transfer_before": "첫 송금 직전", "transfer_seen": "송금 있었나"}
YES, NO = "붙임", "안 붙임"

# 기준표 v2.1 (03_label_guide.md)의 해당 문장
RULES = {
    "1": "3장 요구: 첫 송금 전, 금액이나 보낼 곳을 정한 구체적인 돈 요구. \"보내 주세요\"·\"내 주세요\"처럼 돈을 보내라고 직접 요구하면 금액이 없어도 요구.",
    "2": "3장: 돈도 정보도 아닌 작은 부탁(이메일 대신 보내기 같은 심부름)은 신호 없이 신뢰로 둔다.",
    "3": "D3: 상대 사정을 위해 돈을 내 달라, 빌려 달라, 상대 대신 돈을 보내 달라. 이자·수수료·사례를 주겠다고 해도 D3 (그 약속에는 P1을 함께). / D2: 수익을 내려고 돈을 쓰라.",
    "4": "1장: [사진]·[파일]만 있는 메시지는 바로 앞에 같은 사람이 그것을 설명했으면 그 설명의 신호와 단계를 따른다. 설명이 없으면 신호는 비우고, 단계는 같은 사람의 바로 앞 메시지와 같게.",
    "5": "3장: 단체방 다른 회원의 맞장구·인증·참여 의사는 신뢰 (운영자 메시지만 돈의 흐름으로 봄).",
    "6": "3장 접근: 사례 전체에서 처음 만나는 상대의 첫 메시지와 바로 이어지는 인사·자기소개. 이미 관계가 시작된 뒤 새 방이 열리면 그 방의 첫 메시지는 내용대로.",
    "7": "4.2: \"나\"가 처음 돈을 보내겠다고 하거나 보내는 방법·금액을 묻는 말(묻는 말도 포함) 바로 앞의 상대 메시지. 의사 표현 없이 바로 보냈으면 첫 송금 직전과 같음.",
    "9": "4.1: 돈을 쓰라고 처음 권한 메시지. 가장 먼저 나온 것을 고른다. 내 증권 계좌에서 특정 종목을 사라는 안내 포함, 매도 안내만은 제외.",
    "10": "P1: 큰 수익 약속, 손해 없음, 손실 보상 약속, 언제든 뺄 수 있음, 수익률 알림. / 아님: \"날 수도 있다\" 정도, 작은 작업 보수 안내, 다른 사람이 벌었다는 이야기만.",
    "12": "4.3: \"나\"가 처음 실제로 돈을 보낸 시점(\"나\"의 계좌 밖으로 나간 돈) 바로 앞의 상대 메시지. \"해 볼게요\"는 송금 아님. 뒤의 말로 앞선 송금이 분명하면 송금이 있었을 마지막 송금 지시 번호. 송금이 끝까지 없으면 금액·계좌·주소·사이트 중 하나라도 정해 준 첫 구체 송금 지시.",
    "13": "4.3 송금 있었나: \"나\"의 말(\"송금했어요\") 또는 뒤의 상대 말(\"입금 확인됐습니다\")로 \"나\"의 돈이 계좌 밖으로 나간 것이 분명하면 Y, 끝까지 없으면 N. 뒤의 말로 앞선 송금이 분명한 경우도 Y.",
    "14": "3장 단계는 돈의 흐름 기준. 유도 = 첫 송금 전, 돈을 쓰는 쪽으로 이끎(아직 금액·보낼 곳 없음) / 요구 = 첫 송금 전, 금액이나 보낼 곳을 정한 돈 요구 / 추가요구 = 첫 송금 뒤의 요구·출금 방해·붙잡기·정보 요구 / 신뢰 = 친밀감·권위·증거 (첫 송금 뒤의 안부·사진도 신뢰) / 분류불가 = 어디에도 맞지 않음.",
    "11": "G4: \"나\"의 거절·망설임·그만두겠다는 말 뒤에 돈 요구·투자를 다시 하라고 하거나 붙잡음. / 아님: 거절 없이 다시 요구, 출금 막기(D5), 돈과 무관한 부탁의 반복.",
}


def load_models(folder, models=MODELS):
    msgs, pts = {}, {}
    for m in models:
        msgs[m] = {(r["case"], int(r["seq"])): r for r in read_csv(os.path.join(folder, f"labels_{m}_messages.csv"))}
        pts[m] = {r["case"]: r for r in read_csv(os.path.join(folder, f"labels_{m}_points.csv"))}
    return msgs, pts


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--shuffle-seed", type=int, required=True)
    ap.add_argument("--items", required=True)
    ap.add_argument("--labels", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--models", default=",".join(MODELS), help="쉼표로 구분한 모델 3개")
    ap.add_argument("--guide", default="기준표 v2.1", help="안내·머리글에 쓸 기준표 이름")
    ap.add_argument("--blind-name", default="blind_cases_seed20261021.md", help="금전 지점 판정 때 볼 가림판 파일 이름")
    args = ap.parse_args()
    models = tuple(m.strip() for m in args.models.split(",") if m.strip())
    if len(models) != 3:
        sys.exit(f"[오류] 모델은 3개여야 함: {models}")

    cases = {}
    for i, path in enumerate(ordered(args.files, args.shuffle_seed), 1):
        with open(path, encoding="utf-8") as f:
            lines = [json.loads(l) for l in f if l.strip()]
        cases[f"C{i:02d}"] = lines[1:]
    msgs, pts = load_models(args.labels, models)
    items = read_csv(args.items)

    wb = Workbook()
    ws = wb.active
    ws.title = "안내"
    ws.column_dimensions["A"].width = 110
    for i, (t, f) in enumerate([
            (f"최종 라벨 판정 ({args.guide})", TITLE), ("", BASE),
            (f"판정할 칸 {len(items)}개. 모델 셋의 다수결로 정하지 않는 칸만 모았습니다 (셋이 모두 다른 칸, 기준표에서 규칙이 바뀐 칸, 금전 지점 중 갈린 곳).", BASE),
            ("나머지 칸은 세 모델 다수결(3명 중 2명 이상)로 정합니다.", BASE), ("", BASE),
            ("1. '판정' 시트에서 한 줄씩, 메시지와 앞뒤 메시지를 읽습니다.", BASE),
            (f"2. 오른쪽 '{args.guide}' 칸의 문장을 읽고, 그 문장에 맞는 답을 '판정' 칸 드롭다운에서 고릅니다.", BASE),
            ("3. 모델 답은 참고만 합니다. 기준표 문장과 맞지 않으면 다수의 답이라도 따르지 않습니다.", BASE),
            ("4. 기준표로도 정할 수 없으면 아무거나 고르지 말고 메모에 이유를 적습니다 (기준표를 더 고칠 곳).", BASE),
            (f"5. 금전 지점 줄은 가림판({args.blind_name})에서 그 사례 전체를 보고 고릅니다.", BASE),
            ("", BASE),
            ("이번에는 읽으면서 바로 판단하는 것이 아니라, 기준표 문장을 보고 고르는 작업입니다.", BOLD)], 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font, c.alignment = f, Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("판정")
    heads = ["#", "사례", "메시지 번호", "앞 메시지", "메시지", "뒤 메시지", "판정할 것",
             *models, args.guide, "판정", "메모"]
    widths = [4, 6, 7, 34, 42, 30, 16, 10, 10, 10, 46, 12, 24]
    for j, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=1, column=j, value=h)
        style_cell(c, font=BOLD, fill=HEAD_FILL, align=CENTER)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    dv_stage = DataValidation(type="list", formula1='"' + ",".join(s for s, _ in STAGES) + '"', allow_blank=True)
    dv_yn = DataValidation(type="list", formula1=f'"{YES},{NO}"', allow_blank=True)
    ws.add_data_validation(dv_stage)
    ws.add_data_validation(dv_yn)

    def show(m, rn):
        who = "나" if m["is_user"] else display_sender(m, rn[m["room"]])
        return f"{m['seq']} {who}: {m['text']}"

    for i, it in enumerate(items, 2):
        cid, seq, field, rule = it["case"], int(it["seq"]), it["field"], it["rule"]
        if cid not in cases:
            sys.exit(f"[오류] 없는 사례 {cid}")
        if rule not in RULES:
            sys.exit(f"[오류] 규칙 번호 {rule}이 RULES에 없음")
        cm = cases[cid]
        rn = room_names(cm)
        if field.startswith("point:"):
            key = field.split(":", 1)[1]
            if key not in POINT_NAME:
                sys.exit(f"[오류] 금전 지점 이름 이상: {field}")
            what = POINT_NAME[key]
            answers = [pts[m][cid].get(key, "") for m in models]
            prev, nxt = "(사례 전체는 가림판에서)", ""
            msg = "\"나\"의 메시지:\n" + "\n".join(show(m, rn) for m in cm if m["is_user"])
            seq_cell = ""
            opts = ["Y", "N"] if key == "transfer_seen" else [str(m["seq"]) for m in cm if not m["is_user"]]
            dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
            ws.add_data_validation(dv)
        else:
            idx = next((k for k, m in enumerate(cm) if m["seq"] == seq), None)
            if idx is None or cm[idx]["is_user"]:
                sys.exit(f"[오류] {cid} {seq}: 상대방 메시지가 아님")
            prev = "\n".join(show(m, rn) for m in cm[max(0, idx - 2):idx])
            msg = show(cm[idx], rn)
            nxt = show(cm[idx + 1], rn) if idx + 1 < len(cm) else ""
            seq_cell = seq
            if field == "stage":
                what = "단계"
                answers = [msgs[m][(cid, seq)]["stage"] for m in models]
                dv = dv_stage
            elif field.startswith("sig:"):
                code = field.split(":", 1)[1]
                if code not in SIGNAL_NAME:
                    sys.exit(f"[오류] 신호 코드 이상: {field}")
                what = f"{code} {SIGNAL_NAME[code]}"
                answers = [YES if code in parse_signals(msgs[m][(cid, seq)]["signals"]) else NO for m in models]
                dv = dv_yn
            else:
                sys.exit(f"[오류] field 이상: {field}")
        vals = [i - 1, cid, seq_cell, prev, msg, nxt, what, *answers, RULES[rule]]
        for j, v in enumerate(vals, 1):
            style_cell(ws.cell(row=i, column=j, value=v), align=WRAP if j in (4, 5, 6, 11) else CENTER)
        for j in (12, 13):
            c = ws.cell(row=i, column=j)
            style_cell(c, fill=INPUT_FILL, align=CENTER if j == 12 else WRAP)
            c.protection = Protection(locked=False)
        dv.add(ws.cell(row=i, column=12))
        ws.row_dimensions[i].height = max(48, 15 * (1 + max(len(str(prev)), len(str(msg)), len(RULES[rule])) // 22))
    ws.freeze_panes = "H2"
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    wb.save(args.out)
    print(f"판정할 칸 {len(items)}개 → {args.out}")


if __name__ == "__main__":
    main()
