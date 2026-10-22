#!/usr/bin/env python3
"""
사람 라벨러용 엑셀 양식. 입력 대신 고르기(드롭다운)로 라벨을 붙인다.

사용법
  python tools/scam/make_label_xlsx.py data/eval/scam/*.jsonl --out <파일.xlsx> --shuffle-seed N --cases C03,C05

구성
  안내       채울 칸, 고르는 방법, 지어낸 예시
  코드표     신호 17개·단계 6개 한 줄 기준
  C03 …      사례마다 한 시트. 대화 전체("나" 포함, 회색)를 보여 주고
             상대방 메시지 줄에만 신호 17칸(● 고르기)·단계·확신 드롭다운과 메모 칸
  금전 지점  사례마다 첫 금전 권유·첫 송금 의사 직전·첫 송금 직전(번호 고르기)·송금 여부·이유

번호(C01~)는 make_blind.py·make_label_sheet.py와 같은 규칙. 채운 파일은 agreement.py --human 으로 바로 읽는다.
금전 지점 시트의 열은 머리글 이름으로 찾으므로 열 순서를 바꿔도 되지만 머리글 첫 줄은 바꾸지 말 것.
"""
import argparse
import json
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from make_blind import display_sender, room_names  # noqa: E402
from make_label_sheet import ordered  # noqa: E402

# 기준표 03_label_guide.md v2 2장·3장과 같은 순서·이름 (머리글은 줄이지 않고 기준표 이름 그대로)
SIGNALS = [
    ("R1", "단체방 맞장구", "다른 회원이 수익 인증·감사·참여 의사·수령 인증으로 분위기를 띄움 / 아님: 1:1, 운영자 본인 안내, 단순 인사"),
    ("R2", "친밀감 쌓기", "안부·일상·호의·칭찬, 짧은 기간의 사랑·운명 표현"),
    ("D1", "채널 이동 요구", "다른 앱·방·사이트로 옮겨 오라, 다른 사이트에 가입·계정 등록하라 / 아님: 작업·참고용 링크만"),
    ("D2", "투자금 입금 요구", "수익을 내려고 돈을 쓰라 (입금·충전, 내 증권 계좌에서 특정 종목을 사라는 안내 포함, 권유 포함) / 아님: 자랑만, 매도 안내만, 시황 설명"),
    ("D3", "대납·차용 요구", "상대 사정(통관비·배송비·급전)으로 돈을 내 달라, 빌려 달라 / 출금용이면 D5"),
    ("D4", "지정 계좌·경로 통제", "계좌번호·지갑 주소를 주거나, ATM만·창구 금지처럼 보낼 곳·방법을 정함 / 계좌를 주며 보내라 하면 D2와 함께"),
    ("D5", "출금 조건부 추가 입금", "이미 넣은 돈을 빼려면 수수료·보증금·추가 입금을 하라"),
    ("D6", "비밀 유지·고립 요구", "가족·은행에 말하지 마라, 거짓말을 하라 / 아님: 믿는다는 말만(R2)"),
    ("D7", "개인정보 요구·앱 설치", "신분증·주소·전화번호·계정 정보·인증번호를 나에게 달라, 앱을 깔라 / 아님: 상대가 정보를 주는 메시지, 가입하라(D1)"),
    ("P1", "수익 보장·위험 부정", "큰 수익 약속, 손해 없음, 손실 보상 약속, 언제든 뺄 수 있음, 수익률 알림"),
    ("P2", "권위·신분 과시", "특정 직업·지위·유명인·기관을 내세움(처음 보는 상대의 자기소개 포함), 증명서·신분증·사진 / 아님: 회사원 같은 평범한 소개"),
    ("P3", "긴급성·한정성 압박", "오늘까지, 마지막, 몇 명만·회원만, 급한 사정 / 아님: 단순 일정·이벤트 안내"),
    ("P4", "만남·확인 회피", "영상통화·만남·확인을 미루고 핑계를 댐"),
    ("G1", "소액 선보상", "상대가 나에게 작은 돈·선물을 실제로 줬다고 알림 / 아님: 약속·수익률 알림(P1), 다른 회원의 인증(R1)"),
    ("G2", "화제 급전환", "돈과 무관한 대화가 이 메시지에서 돈 얘기로 바뀜 (넘어가는 메시지 하나만)"),
    ("G3", "요구 금액 증가", "앞보다 큰 금액, 대출 권유, 큰 금액일수록 유리한 규칙 (메모에 비교 번호)"),
    ("G4", "요구 반복", "나의 거절·망설임·그만두겠다는 말 뒤에 다시 요구·붙잡음 (메모에 비교 번호) / 거절 없이 다시 요구는 아님"),
]
STAGES = [("접근", "각 방에서 상대의 첫 메시지와 바로 이어지는 인사·자기소개"),
          ("신뢰", "친밀감·권위·증거·소액 선보상·무료 정보로 믿게 만듦"),
          ("유도", "첫 송금 전, 돈을 쓰는 쪽으로 이끎 (정보 요구 포함) — 금액·보낼 곳은 아직 없음"),
          ("요구", "첫 송금 전, 금액이나 보낼 곳을 정한 구체적인 돈 요구"),
          ("추가요구", "첫 송금 뒤의 요구·출금 방해·붙잡기·정보 요구"),
          ("분류불가", "어느 것에도 맞지 않음 (메모에 이유)")]
MARK = "●"

F = "Arial"
BASE = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
GREY_FONT = Font(name=F, size=10, color="808080", italic=True)
TITLE = Font(name=F, size=13, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF9DB")
USER_FILL = PatternFill("solid", fgColor="F2F2F2")
ON_FILL = PatternFill("solid", fgColor="A9D18E")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

TEXT_COL_WIDTH = 58
CHARS_PER_LINE = 30  # 너비 58에서 한글 기준 대략


def line_height(text, per_line=CHARS_PER_LINE):
    lines = sum(max(1, math.ceil(len(p) / per_line)) for p in str(text).split("\n"))
    return max(18, 15 * lines + 4)


def style_cell(c, font=BASE, fill=None, align=WRAP):
    c.font, c.alignment, c.border = font, align, BOX
    if fill:
        c.fill = fill


def guide_sheet(wb, cases):
    ws = wb.active
    ws.title = "안내"
    ws.column_dimensions["A"].width = 100
    rows = [
        ("라벨링 양식 (기준표 v2)", TITLE),
        ("", BASE),
        ("채울 칸: 노란색 칸만 채웁니다. 나머지 칸은 잠겨 있습니다 (검토 → 시트 보호 해제로 풀 수 있음, 암호 없음).", BOLD),
        ("", BASE),
        (f"1. 사례 시트({', '.join(cases)})를 하나씩 엽니다. 대화 전체가 순서대로 있고, \"나\"의 메시지는 회색입니다.", BASE),
        ("2. 먼저 사례를 처음부터 끝까지 한 번 읽습니다. \"나\"가 처음 보내겠다고 밝힌 곳과 처음 실제로 돈을 보낸 곳을 기억해 둡니다.", BASE),
        (f"3. 상대방 메시지 줄에서, 해당하는 신호 칸을 눌러 드롭다운에서 {MARK}를 고릅니다 (여러 개 가능, 없으면 비움). 고르면 칸이 초록색이 됩니다. 지우려면 Delete.", BASE),
        ("4. 단계·확신 칸은 드롭다운에서 하나씩 고릅니다. 상대방 메시지마다 둘 다 채웁니다 (오른쪽 '확인' 칸에 빠진 것이 표시됨).", BASE),
        ("5. 메모: G3·G4를 골랐으면 비교한 앞 메시지 번호, 고민한 다른 신호, 기준표에 없는 행동(\"기준표에 없음: ○○\")을 적습니다.", BASE),
        ("6. 사례를 마칠 때마다 '금전 지점' 시트의 그 줄을 채웁니다. 번호는 드롭다운에서 고릅니다.", BASE),
        ("7. 다 끝나면 그대로 저장(.xlsx)해서 가져오면 됩니다. CSV로 바꿀 필요 없습니다.", BASE),
        ("", BASE),
        ("헷갈리기 쉬운 짝", BOLD),
        ("· D2 / D3 / D5: 왜 돈을 내라는가 — 벌려고(D2), 상대 사정으로(D3), 빼려고(D5)", BASE),
        ("· P1 / G1: 주겠다는 약속·수익률 알림이면 P1, 상대가 나에게 실제로 줬다면 G1 (다른 회원의 인증은 R1)", BASE),
        ("· D1 / D7: 가입·계정을 만들라면 D1, 신분증·계정 정보를 달라면 D7. 상대가 정보를 주는 메시지(시황·종목 안내)는 D7이 아님", BASE),
        ("· D2: 내 증권 계좌에서 특정 종목을 사라는 안내도 D2 (팔라는 안내만 있으면 아님). 계좌·지갑을 주며 보내라 하면 D4도 함께", BASE),
        ("· G4: 내가 거절·망설인 뒤에만. 거절 없이 다시 요구하는 것은 G4가 아님", BASE),
        ("· \"저 40% 벌었어요\" 같은 자랑은 D2가 아니라 P1 (단체방이면 R1도)", BASE),
        ("· 단계: 각 방의 첫 메시지는 접근. 첫 송금 전에는 막연한 권유면 유도, 금액·보낼 곳을 정했으면 요구. 첫 송금 뒤의 요구는 모두 추가요구", BASE),
        ("· 판단이 어려우면 확신을 '낮음'으로 두고 메모에 고민한 내용을 적습니다. 오래 고민하지 않아도 됩니다.", BASE),
        ("", BASE),
        ("예시 (지어낸 대화 — 평가 사례 문장이 아님)", BOLD),
        ("· 4 \"오늘 점심은 드셨어요?\" → R2 ● / 단계 신뢰 / 확신 높음", BASE),
        ("· 7 \"참, 요즘 코인 해 보셨어요? 저는 이걸로 꽤 벌었어요\" → G2 ●, P1 ● / 유도 / 높음 / 메모: 권유는 아님(자랑)", BASE),
        ("· 9 \"처음엔 50만 원만 이 사이트에 넣어 보세요 (링크)\" → D2 ●, D4 ● / 요구 / 낮음 / 메모: D1도 고민", BASE),
        ("", BASE),
        ("원래 사례 파일, 분석 문서, 모델 라벨 결과는 라벨링을 마칠 때까지 열지 않습니다.", BOLD),
    ]
    for i, (t, f) in enumerate(rows, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font, c.alignment = f, Alignment(wrap_text=True, vertical="top")


def code_sheet(wb):
    ws = wb.create_sheet("코드표")
    for col, w in zip("ABC", (8, 22, 100)):
        ws.column_dimensions[col].width = w
    r = 1
    for title, items in (("신호 (여러 개 가능)", [(a, b, c) for a, b, c in SIGNALS]),
                         ("단계 (하나만)", [(a, "", b) for a, b in STAGES])):
        ws.cell(row=r, column=1, value=title).font = BOLD
        r += 1
        for vals in items:
            for j, v in enumerate(vals, 1):
                style_cell(ws.cell(row=r, column=j, value=v), font=BOLD if j == 1 else BASE)
            r += 1
        r += 1
    ws.cell(row=r, column=1, value="자세한 기준과 '붙이지 않는 경우'는 03_label_guide.md 2~4장").font = GREY_FONT
    ws.protection.sheet = True


def case_sheet(wb, cid, msgs, dv_mark, dv_stage, dv_conf):
    ws = wb.create_sheet(cid)
    rn = room_names(msgs)
    heads = ["#", "날짜", "방", "보낸 사람", "메시지"] + [f"{c}\n{n}" for c, n, _ in SIGNALS] + ["단계", "확신", "메모", "확인"]
    widths = [5, 6, 11, 12, TEXT_COL_WIDTH] + [9] * len(SIGNALS) + [10, 7, 30, 14]
    for j, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=1, column=j, value=h)
        style_cell(c, font=BOLD, fill=HEAD_FILL, align=CENTER)
        ws.column_dimensions[c.column_letter].width = w
    for j, (code, name, desc) in enumerate(SIGNALS, 6):
        cm = Comment(f"{code} {name}\n{desc}", "기준표")
        cm.width, cm.height = 320, 130
        ws.cell(row=1, column=j).comment = cm
    ws.row_dimensions[1].height = 62
    s0 = 6
    s1 = s0 + len(SIGNALS) - 1
    c_stage, c_conf, c_note, c_check = s1 + 1, s1 + 2, s1 + 3, s1 + 4
    L = lambda j: ws.cell(row=1, column=j).column_letter  # noqa: E731

    for i, m in enumerate(msgs, 2):
        user = m["is_user"]
        vals = [m["seq"], f"d{m['day']}", rn[m["room"]], display_sender(m, rn[m["room"]]), m["text"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            style_cell(c, font=GREY_FONT if user else BASE, fill=USER_FILL if user else None,
                       align=WRAP if j == 5 else Alignment(vertical="top", horizontal="center", wrap_text=True))
        ws.row_dimensions[i].height = line_height(m["text"])
        for j in range(s0, c_check + 1):
            c = ws.cell(row=i, column=j)
            if user:
                style_cell(c, fill=USER_FILL)
                continue
            if j == c_check:
                style_cell(c, font=Font(name=F, size=9, color="C00000"))
                continue
            style_cell(c, fill=INPUT_FILL, align=WRAP if j == c_note else CENTER)
            c.protection = Protection(locked=False)
            if s0 <= j <= s1:
                dv_mark.add(c)
            elif j == c_stage:
                dv_stage.add(c)
            elif j == c_conf:
                dv_conf.add(c)
        if not user:
            st, cf = f"{L(c_stage)}{i}", f"{L(c_conf)}{i}"
            ws.cell(row=i, column=c_check,
                    value=f'=IF(AND({st}="",{cf}=""),"단계·확신 비어 있음",IF({st}="","단계 비어 있음",IF({cf}="","확신 비어 있음","")))')

    last = len(msgs) + 1
    rng = f"{L(s0)}2:{L(s1)}{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{L(s0)}2="{MARK}"'], fill=ON_FILL))
    ws.freeze_panes = "F2"
    ws.protection.sheet = True
    ws.protection.formatColumns = False  # 열 너비·행 높이 조절은 허용
    ws.protection.formatRows = False
    return [m["seq"] for m in msgs if not m["is_user"]]


def points_sheet(wb, cases):
    ws = wb.create_sheet("금전 지점")
    heads = ["사례", "메시지 수", "첫 금전 권유\n(번호)", "첫 송금 의사 직전\n(번호)", "첫 송금 직전\n(번호)", "송금 있었나\n(Y/N)",
             "고른 이유 (세 지점 각각, 고민한 다른 번호도)"]
    for j, (h, w) in enumerate(zip(heads, (8, 9, 13, 15, 13, 12, 70)), 1):
        c = ws.cell(row=1, column=j, value=h)
        style_cell(c, font=BOLD, fill=HEAD_FILL, align=CENTER)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 34
    dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True,
                           errorTitle="Y 또는 N", error="송금이 있었으면 Y, 없었으면 N을 고르세요.")
    ws.add_data_validation(dv_yn)
    for i, (cid, n, seqs) in enumerate(cases, 2):
        for j, v in enumerate((cid, n), 1):
            style_cell(ws.cell(row=i, column=j, value=v), align=CENTER)
        lst = ",".join(str(s) for s in seqs)
        dv = DataValidation(type="list", formula1=f'"{lst}"', allow_blank=True,
                            errorTitle="상대방 메시지 번호", error="이 사례의 상대방 메시지 번호 중에서 고르세요.")
        ws.add_data_validation(dv)
        for j in (3, 4, 5, 6, 7):
            c = ws.cell(row=i, column=j)
            style_cell(c, fill=INPUT_FILL, align=WRAP if j == 7 else CENTER)
            c.protection = Protection(locked=False)
        for j in (3, 4, 5):
            dv.add(ws.cell(row=i, column=j))
        dv_yn.add(ws.cell(row=i, column=6))
        ws.row_dimensions[i].height = 48
    note = len(cases) + 3
    for k, t in enumerate((
            "첫 금전 권유 (기준표 4.1): 상대가 돈을 쓰라고 처음 권한 메시지. 구체 지시가 없어도 되고, 내 증권 계좌에서 특정 종목을 사라는 안내도 포함.",
            "첫 송금 의사 직전 (4.2): 내가 처음 보내겠다고 하거나 보내는 방법을 물은 메시지 바로 앞의 상대 메시지. 의사 표현 없이 바로 보냈으면 첫 송금 직전과 같은 번호.",
            "첫 송금 직전 (4.3): 내 돈이 내 계좌 밖으로 처음 나가기 바로 앞의 상대 메시지 (다른 방이어도 됨). 내 증권 계좌 안의 매매는 송금이 아님. \"해 볼게요\"는 의사 표현이지 송금이 아님.",
            "뒤의 말로 앞선 송금이 분명하면 그 송금이 있었을 마지막 송금 지시. 송금이 끝까지 없으면 금액·계좌·사이트 중 하나라도 처음 정해 준 메시지 (금액만 있어도 됨), 송금 칸은 N.",
            "\"보냈어요\"가 돈인지 꼭 확인 — 이메일·사진·캡처는 송금이 아님. 번호 칸은 상대방 메시지 번호만 고를 수 있고, 같은 번호여도 됩니다.")):
        c = ws.cell(row=note + k, column=1, value=t)
        c.font = BASE
        ws.merge_cells(start_row=note + k, start_column=1, end_row=note + k, end_column=7)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[note + k].height = 30
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--out", required=True)
    ap.add_argument("--shuffle-seed", type=int, default=None)
    ap.add_argument("--cases", required=True, help="쉼표로 구분한 가림판 번호 (예: C03,C07)")
    args = ap.parse_args()
    if not args.out.endswith(".xlsx"):
        sys.exit("[오류] --out은 .xlsx 파일 경로")
    want = [c.strip().upper() for c in args.cases.split(",") if c.strip()]
    if len(set(want)) != len(want):
        sys.exit("[오류] --cases에 같은 번호가 두 번 있음")

    by_id = {}
    for i, path in enumerate(ordered(args.files, args.shuffle_seed), 1):
        by_id[f"C{i:02d}"] = path
    missing = [c for c in want if c not in by_id]
    if missing:
        sys.exit(f"[오류] 없는 사례 번호: {missing}")

    wb = Workbook()
    guide_sheet(wb, want)
    code_sheet(wb)
    cases = []
    for cid in want:
        with open(by_id[cid], encoding="utf-8") as f:
            lines = [json.loads(l) for l in f if l.strip()]
        if lines[0].get("type") != "meta":
            sys.exit(f"[오류] {by_id[cid]}: 첫 줄이 meta가 아님")
        msgs = lines[1:]
        dv_mark = DataValidation(type="list", formula1=f'"{MARK}"', allow_blank=True,
                                 errorTitle="고르기만 가능", error=f"드롭다운에서 {MARK}를 고르거나 Delete로 지우세요.")
        dv_stage = DataValidation(type="list", formula1='"' + ",".join(s for s, _ in STAGES) + '"', allow_blank=True,
                                  errorTitle="단계", error="드롭다운에서 단계를 고르세요.")
        dv_conf = DataValidation(type="list", formula1='"높음,낮음"', allow_blank=True,
                                 errorTitle="확신", error="높음 또는 낮음을 고르세요.")
        seqs = case_sheet(wb, cid, msgs, dv_mark, dv_stage, dv_conf)
        ws = wb[cid]
        for dv in (dv_mark, dv_stage, dv_conf):
            ws.add_data_validation(dv)
        cases.append((cid, len(msgs), seqs))
    points_sheet(wb, cases)
    wb.calculation.fullCalcOnLoad = True  # '확인' 칸 수식을 열 때 바로 계산
    wb.save(args.out)
    n = sum(len(s) for _, _, s in cases)
    print(f"사례 {len(cases)}건, 상대방 메시지 {n}줄 → {args.out}")


if __name__ == "__main__":
    main()
