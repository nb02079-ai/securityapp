#!/usr/bin/env python3
"""
최종 라벨(평가 정답지) 만들기: 모델 3개 다수결 + 사람 판정(adjudication)을 합친다.

사용법
  python tools/scam/finalize_labels.py data/eval/scam/*.jsonl --shuffle-seed 20261021 \
      --labels docs/labeling/label_results/v2 \
      --adjudication <판정을 채운 adjudication .xlsx> \
      --out data/eval/scam/labels
  홀드아웃: --models chatgpt,gemini,grok (기본 grok,claude,chatgpt)

규칙 (docs/labeling/06_label_round2.md 4장)
  - 신호: 메시지 × 신호마다 세 모델 중 2명 이상이 붙였으면 붙임. 판정한 칸은 판정을 따름
  - 단계: 세 모델 중 2명 이상이 같은 단계. 판정한 칸은 판정을 따름. 다수결도 판정도 없으면 멈춤
  - 금전 지점: 세 모델이 같으면 그 값. 판정한 지점은 판정을 따름. 갈렸는데 판정이 없으면 멈춤

결과 (사례 이름은 가림판 번호가 아니라 원래 case_id)
  final_messages.csv  case_id, seq, signals, stage, signals_source, stage_source
  final_points.csv    case_id, first_suggest, first_intent_before, first_transfer_before, transfer_seen, source
"""
import argparse
import csv
import json
import os
import sys
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from agreement import CODES, STAGE_NAMES, parse_signals, read_csv  # noqa: E402
from make_adjudication_xlsx import MODELS, NO, POINT_NAME, YES, load_models  # noqa: E402
from make_label_sheet import ordered  # noqa: E402

POINT_KEYS = ("first_suggest", "first_intent_before", "first_transfer_before", "transfer_seen")
NAME_TO_POINT = {v: k for k, v in POINT_NAME.items()}


def load_adjudication(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["판정"]
    head = [c.value for c in ws[1]]
    i_case, i_seq, i_what, i_dec = head.index("사례"), head.index("메시지 번호"), head.index("판정할 것"), head.index("판정")
    adj = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        cid, seq, what, dec = row[i_case], row[i_seq], row[i_what], row[i_dec]
        if dec in (None, ""):
            sys.exit(f"[오류] 판정이 비어 있음: {cid} {seq} {what}")
        dec = str(dec).strip()
        if what in NAME_TO_POINT:
            key = ("point", cid, NAME_TO_POINT[what])
            if key[2] == "transfer_seen":
                dec = dec.upper()
                if dec not in ("Y", "N"):
                    sys.exit(f"[오류] 송금 있었나 판정은 Y/N: {cid} '{dec}'")
            elif not dec.isdigit():
                sys.exit(f"[오류] 금전 지점 판정이 번호가 아님: {cid} {what} '{dec}'")
        elif what == "단계":
            key = ("stage", cid, int(seq))
            if dec not in STAGE_NAMES:
                sys.exit(f"[오류] 단계 판정 값 이상: {cid} {seq} '{dec}'")
        else:
            code = str(what).split()[0]
            if code not in CODES or dec not in (YES, NO):
                sys.exit(f"[오류] 신호 판정 이상: {cid} {seq} {what} '{dec}'")
            key = ("sig", cid, int(seq), code)
        if key in adj:
            sys.exit(f"[오류] 같은 칸을 두 번 판정함: {key}")
        adj[key] = dec
    return adj


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--shuffle-seed", type=int, required=True)
    ap.add_argument("--labels", required=True)
    ap.add_argument("--adjudication", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--models", default=",".join(MODELS), help="쉼표로 구분한 모델 3개")
    args = ap.parse_args()
    models = tuple(m.strip() for m in args.models.split(",") if m.strip())
    if len(models) != 3:
        sys.exit(f"[오류] 모델은 3개여야 함: {models}")

    blind = {}  # 가림판 번호 → (case_id, 메시지들)
    for i, path in enumerate(ordered(args.files, args.shuffle_seed), 1):
        with open(path, encoding="utf-8") as f:
            lines = [json.loads(l) for l in f if l.strip()]
        blind[f"C{i:02d}"] = (lines[0]["case_id"], lines[1:])
    msgs, pts = load_models(args.labels, models)
    adj = load_adjudication(args.adjudication)
    used = set()

    out_msgs, n_adj_sig, n_adj_stage = [], 0, 0
    for cid in sorted(blind):
        case_id, cm = blind[cid]
        for m in cm:
            if m["is_user"]:
                continue
            k = (cid, m["seq"])
            for mod in models:
                if k not in msgs[mod]:
                    sys.exit(f"[오류] {mod}에 {k} 없음")
                if msgs[mod][k]["text"] != m["text"]:
                    sys.exit(f"[오류] {mod} {k} 메시지 내용이 평가 세트와 다름 (번호가 어긋났을 수 있음)")
            votes = Counter(c for mod in models for c in parse_signals(msgs[mod][k]["signals"]) if c in CODES)
            sig, sig_src = set(), "다수결"
            for code in CODES:
                key = ("sig", cid, m["seq"], code)
                if key in adj:
                    used.add(key)
                    n_adj_sig += 1
                    sig_src = "다수결+판정"
                    if adj[key] == YES:
                        sig.add(code)
                elif votes[code] >= 2:
                    sig.add(code)
            key = ("stage", cid, m["seq"])
            if key in adj:
                used.add(key)
                n_adj_stage += 1
                stage, st_src = adj[key], "판정"
            else:
                top, cnt = Counter(msgs[mod][k]["stage"] for mod in models).most_common(1)[0]
                if cnt < 2:
                    sys.exit(f"[오류] {k} 단계 다수결도 판정도 없음")
                stage, st_src = top, "다수결"
            out_msgs.append({"case_id": case_id, "seq": m["seq"], "signals": ";".join(c for c in CODES if c in sig),
                             "stage": stage, "signals_source": sig_src, "stage_source": st_src})

    out_pts = []
    for cid in sorted(blind):
        case_id, cm = blind[cid]
        row, src = {"case_id": case_id}, []
        their = {m["seq"] for m in cm if not m["is_user"]}
        for f in POINT_KEYS:
            key = ("point", cid, f)
            if key in adj:
                used.add(key)
                row[f] = adj[key]
                src.append(f"{f}=판정")
                continue
            vals = {str(pts[mod][cid].get(f, "")).strip().upper() for mod in models}
            if len(vals) != 1 or "" in vals:
                sys.exit(f"[오류] {cid} {f}: 모델 값이 갈렸는데 판정이 없음 {vals}")
            row[f] = vals.pop()
        for f in POINT_KEYS[:3]:
            if int(row[f]) not in their:
                sys.exit(f"[오류] {cid} {f}={row[f]}: 상대방 메시지 번호가 아님")
        if row["transfer_seen"] not in ("Y", "N"):
            sys.exit(f"[오류] {cid} transfer_seen 값 이상")
        row["source"] = "세 모델 일치" if not src else "; ".join(src)
        out_pts.append(row)

    unused = set(adj) - used
    if unused:
        sys.exit(f"[오류] 쓰이지 않은 판정 (번호가 어긋났을 수 있음): {sorted(unused)}")

    os.makedirs(args.out, exist_ok=True)
    for name, rows in (("final_messages.csv", out_msgs), ("final_points.csv", out_pts)):
        with open(os.path.join(args.out, name), "w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            w.writeheader()
            w.writerows(rows)
    print(f"메시지 {len(out_msgs)}줄 (신호 판정 {n_adj_sig}칸, 단계 판정 {n_adj_stage}칸), 사례 {len(out_pts)}건 → {args.out}")


if __name__ == "__main__":
    main()
