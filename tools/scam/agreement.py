#!/usr/bin/env python3
"""
라벨러 간 일치율 계산 + 다수결 라벨 만들기 (plan.md 6.1 보조 지표 "라벨러 간 일치율").

사용법
  python tools/scam/agreement.py \
      --sheet <make_label_sheet.py로 만든 label_messages.csv (10건 전체)> \
      --model grok=<msg.csv>,<points.csv> --model claude=... --model chatgpt=... \
      [--human me=<make_label_xlsx.py 양식을 채운 .xlsx>] \
      --out <저장소 밖 폴더>

결과
  report.md               점검 결과, 신호별 κ, 단계 κ, 금전 지점 표
  majority_messages.csv   모델 다수결 신호(3명 중 2명 이상)·단계, 갈린 칸 표시
  disagreements.csv       신호·단계가 갈린 메시지 목록 (기준표를 고칠 곳 찾기용)

계산 방식
  - 신호: 메시지 × 신호마다 있음/없음 두 값. 모델 3명은 Fleiss κ, 둘씩은 Cohen κ,
    사람은 같은 메시지에서 모델 다수결과 Cohen κ. 한쪽이라도 한 번도 안 붙인 신호는 κ를 계산하지 않음(-).
  - 단계: 메시지마다 하나. 같은 방식.
  - 금전 지점: 사례마다 번호가 같은지, ±1(상대 메시지 기준 바로 옆)까지 같은지. v1 결과(의사 지점 없음)도 읽음.
"""
import argparse
import csv
import os
import sys
from collections import Counter
from itertools import combinations

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from make_label_xlsx import SIGNALS, STAGES, MARK  # noqa: E402

CODES = [c for c, _, _ in SIGNALS]
STAGE_NAMES = [s for s, _ in STAGES]


def read_csv(path):
    for enc in ("utf-8-sig", "cp949"):
        try:
            with open(path, encoding=enc, newline="") as f:
                return list(csv.DictReader(f))
        except UnicodeDecodeError:
            continue
    sys.exit(f"[오류] 인코딩을 읽을 수 없음: {path}")


def parse_signals(s):
    return {x.strip().upper() for x in (s or "").replace(",", ";").split(";") if x.strip()}


def load_model(name, msg_path, pts_path, sheet_keys, problems):
    rows = read_csv(msg_path)
    labels = {}
    for r in rows:
        key = (r["case"].strip(), int(r["seq"]))
        if key not in sheet_keys:
            problems.append(f"{name}: 시트에 없는 줄 {key}")
            continue
        if r["text"] != sheet_keys[key]:
            problems.append(f"{name}: {key} 메시지 내용이 시트와 다름")
        sig = parse_signals(r["signals"])
        bad = sig - set(CODES)
        if bad:
            problems.append(f"{name}: {key} 기준표에 없는 코드 {sorted(bad)}")
        stage = (r["stage"] or "").strip()
        if stage not in STAGE_NAMES:
            problems.append(f"{name}: {key} 단계 값 이상 '{stage}'")
        if key in labels:
            problems.append(f"{name}: {key} 중복")
        labels[key] = {"signals": sig & set(CODES), "stage": stage, "conf": (r.get("confidence") or "").strip(),
                       "note": r.get("note", "")}
    missing = set(sheet_keys) - set(labels)
    if missing:
        problems.append(f"{name}: 빠진 줄 {len(missing)}개 (예: {sorted(missing)[:3]})")
    pts = {r["case"].strip(): r for r in read_csv(pts_path)}
    return labels, pts


def load_human(path, sheet_keys, problems):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    labels, pts = {}, {}
    for ws in wb.worksheets:
        if not (ws.title.startswith("C") and ws.title[1:].isdigit()):
            continue
        head = [c.value for c in ws[1]]
        sig_cols = {i: str(h).split("\n")[0] for i, h in enumerate(head) if h and str(h).split("\n")[0] in CODES}
        i_stage, i_conf, i_note = head.index("단계"), head.index("확신"), head.index("메모")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None or row[3] == "나":
                continue
            key = (ws.title, int(row[0]))
            if key not in sheet_keys:
                problems.append(f"사람: 시트에 없는 줄 {key}")
                continue
            if row[4] != sheet_keys[key]:
                problems.append(f"사람: {key} 메시지 내용이 시트와 다름")
            sig = {code for i, code in sig_cols.items() if row[i] == MARK}
            stage = (row[i_stage] or "").strip()
            if stage not in STAGE_NAMES:
                problems.append(f"사람: {key} 단계 값 이상 '{stage}'")
            labels[key] = {"signals": sig, "stage": stage, "conf": row[i_conf] or "", "note": row[i_note] or ""}
    ws = wb["금전 지점"]
    # 머리글 첫 줄로 열을 찾는다 (v1 양식에는 "첫 송금 의사 직전" 열이 없음)
    head = [str(c.value).split("\n")[0].strip() if c.value else "" for c in ws[1]]
    col = {"first_suggest": "첫 금전 권유", "first_intent_before": "첫 송금 의사 직전",
           "first_transfer_before": "첫 송금 직전", "transfer_seen": "송금 있었나", "reason": "고른 이유"}
    idx = {k: head.index(v) for k, v in col.items() if v in head}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and str(row[0]).startswith("C"):
            pts[row[0]] = {k: row[i] for k, i in idx.items()}
    cases = {k[0] for k in labels}
    expect = {k for k in sheet_keys if k[0] in cases}
    if expect - set(labels):
        problems.append(f"사람: 빠진 줄 {len(expect - set(labels))}개")
    return labels, pts


def cohen(a, b):
    n = len(a)
    if n == 0:
        return None
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / (n * n)
    return None if pe == 1 else (po - pe) / (1 - pe)


def fleiss(table):
    """table: 항목마다 [범주별 평가자 수]. 평가자 수는 모든 항목에서 같아야 함."""
    N = len(table)
    if N == 0:
        return None
    n = sum(table[0])
    if n < 2 or any(sum(row) != n for row in table):
        return None  # 빈 단계 값 등으로 평가자 수가 맞지 않으면 계산하지 않음 (점검 문제에 따로 나옴)
    k = len(table[0])
    p = [sum(row[j] for row in table) / (N * n) for j in range(k)]
    P = [(sum(c * c for c in row) - n) / (n * (n - 1)) for row in table]
    Pbar, Pe = sum(P) / N, sum(x * x for x in p)
    return None if Pe == 1 else (Pbar - Pe) / (1 - Pe)


def fmt(x):
    return "-" if x is None else f"{x:.2f}"


def int_or_none(v):
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--model", action="append", required=True, help="이름=메시지.csv,지점.csv")
    ap.add_argument("--human", default=None, help="이름=양식.xlsx")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)

    sheet = read_csv(args.sheet)
    sheet_keys = {(r["case"], int(r["seq"])): r["text"] for r in sheet}
    order = sorted(sheet_keys, key=lambda k: (k[0], k[1]))
    problems = []

    models = {}
    for spec in args.model:
        name, paths = spec.split("=", 1)
        mp, pp = paths.split(",")
        models[name] = load_model(name, mp, pp, sheet_keys, problems)
    names = list(models)
    if len(names) != 3:
        sys.exit("[오류] 다수결·Fleiss κ는 모델 3개 기준")
    human = None
    if args.human:
        hname, hpath = args.human.split("=", 1)
        human = (hname, *load_human(hpath, sheet_keys, problems))

    common = [k for k in order if all(k in models[n][0] for n in names)]

    # 다수결
    maj = {}
    for k in common:
        votes = Counter(c for n in names for c in models[n][0][k]["signals"])
        sig = {c for c, v in votes.items() if v >= 2}
        split_sig = sorted(c for c, v in votes.items() if v == 1 or (v == 2))
        st = Counter(models[n][0][k]["stage"] for n in names)
        top, cnt = st.most_common(1)[0]
        maj[k] = {"signals": sig, "stage": top if cnt >= 2 else "", "split_sig": split_sig,
                  "stage_votes": st}

    out = []
    out.append("# 라벨러 간 일치율\n")
    out.append(f"- 시트 상대방 메시지 {len(sheet_keys)}줄, 모델 {', '.join(names)}"
               + (f", 사람 {human[0]} ({len(human[1])}줄)" if human else ""))
    out.append(f"- 점검 문제 {len(problems)}개" + ("" if not problems else ":"))
    for p in problems:
        out.append(f"  - {p}")

    # 신호별 κ
    out.append("\n## 1. 신호별 일치 (메시지 × 신호의 있음·없음)\n")
    head = "| 코드 | 붙인 수 (" + " / ".join(names) + ") | Fleiss κ (모델 3) | " + " | ".join(
        f"{a}–{b}" for a, b in combinations(names, 2)) + " |"
    if human:
        head += f" {human[0]} 붙인 수 | {human[0]}–다수결 κ |"
    out.append(head)
    out.append("|" + "---|" * (head.count("|") - 1))
    all_tables, hum_a, hum_b = [], [], []
    hkeys = [k for k in common if human and k in human[1]]
    for code in CODES:
        vec = {n: [code in models[n][0][k]["signals"] for k in common] for n in names}
        counts = [sum(vec[n]) for n in names]
        table = [[sum(not vec[n][i] for n in names), sum(vec[n][i] for n in names)] for i in range(len(common))]
        all_tables += table
        fk = fleiss(table) if all(counts) else None
        pair = [cohen(vec[a], vec[b]) if counts[names.index(a)] and counts[names.index(b)] else None
                for a, b in combinations(names, 2)]
        line = f"| {code} | {' / '.join(map(str, counts))} | {fmt(fk)} | " + " | ".join(map(fmt, pair)) + " |"
        if human:
            ha = [code in human[1][k]["signals"] for k in hkeys]
            hb = [code in maj[k]["signals"] for k in hkeys]
            hum_a += ha
            hum_b += hb
            hk = cohen(ha, hb) if any(ha) and any(hb) else None
            line += f" {sum(ha)} | {fmt(hk)} |"
        out.append(line)
    out.append(f"\n- 17개 신호를 모두 합친 Fleiss κ (모델 3): **{fmt(fleiss(all_tables))}**")
    if human:
        out.append(f"- 17개 신호를 모두 합친 {human[0]}–다수결 Cohen κ ({len(hkeys)}줄): **{fmt(cohen(hum_a, hum_b))}**")
    exact = sum(len({frozenset(models[n][0][k]['signals']) for n in names}) == 1 for k in common)
    out.append(f"- 세 모델의 신호 묶음이 완전히 같은 메시지: {exact}/{len(common)}")

    # 단계 κ
    out.append("\n## 2. 단계 일치\n")
    cat = STAGE_NAMES
    table = [[sum(models[n][0][k]["stage"] == c for n in names) for c in cat] for k in common]
    out.append(f"- Fleiss κ (모델 3): **{fmt(fleiss(table))}**")
    for a, b in combinations(names, 2):
        out.append(f"- {a}–{b} Cohen κ: {fmt(cohen([models[a][0][k]['stage'] for k in common], [models[b][0][k]['stage'] for k in common]))}")
    if human:
        hk = [k for k in hkeys if maj[k]["stage"]]
        out.append(f"- {human[0]}–다수결 Cohen κ ({len(hk)}줄, 다수결이 없는 줄 제외): "
                   f"**{fmt(cohen([human[1][k]['stage'] for k in hk], [maj[k]['stage'] for k in hk]))}**")
    out.append("\n| 단계 | " + " | ".join(names) + (f" | {human[0]} (5건)" if human else "") + " |")
    out.append("|---|" + "---|" * (len(names) + (1 if human else 0)))
    for c in cat:
        line = f"| {c} | " + " | ".join(str(sum(models[n][0][k]['stage'] == c for k in common)) for n in names)
        if human:
            line += f" | {sum(human[1][k]['stage'] == c for k in hkeys)}"
        out.append(line + " |")

    # 금전 지점
    out.append("\n## 3. 금전 지점\n")
    raters = [(n, models[n][1]) for n in names] + ([(human[0], human[2])] if human else [])
    out.append("| 사례 | 지점 | " + " | ".join(r for r, _ in raters) + " | 모델 3 일치 |")
    out.append("|---|---|" + "---|" * (len(raters) + 1))
    cases = sorted({k[0] for k in sheet_keys})
    seqs = {c: sorted(s for (cc, s) in sheet_keys if cc == c) for c in cases}
    agree = Counter()
    for c in cases:
        for field, label in (("first_suggest", "첫 금전 권유"), ("first_intent_before", "첫 송금 의사 직전"),
                             ("first_transfer_before", "첫 송금 직전"), ("transfer_seen", "송금 있었나")):
            if not any(field in (p.get(c) or {}) for _, p in raters):
                continue  # v1 결과에는 의사 지점이 없음
            vals = [(p.get(c) or {}).get(field) for _, p in raters]
            mv = vals[:3]
            if field == "transfer_seen":
                same = len({str(v).strip().upper() for v in mv}) == 1
                mark = "일치" if same else "**갈림**"
            else:
                nums = [int_or_none(v) for v in mv]
                if None in nums:
                    mark = "값 없음"
                elif len(set(nums)) == 1:
                    mark = "일치"
                    agree[field] += 1
                else:
                    # ±1: 상대방 메시지 순서에서 바로 옆
                    idx = [seqs[c].index(x) if x in seqs[c] else None for x in nums]
                    near = None not in idx and max(idx) - min(idx) <= 1
                    mark = "±1" if near else "**갈림**"
                    if near:
                        agree[field + "_near"] += 1
            out.append(f"| {c} | {label} | " + " | ".join("" if v is None else str(v) for v in vals) + f" | {mark} |")
    out.append(f"\n- 첫 금전 권유: 세 모델 완전 일치 {agree['first_suggest']}/{len(cases)}, ±1 포함 {agree['first_suggest'] + agree['first_suggest_near']}/{len(cases)}")
    if any("first_intent_before" in (p.get(c) or {}) for _, p in raters for c in cases):
        out.append(f"- 첫 송금 의사 직전: 세 모델 완전 일치 {agree['first_intent_before']}/{len(cases)}, ±1 포함 {agree['first_intent_before'] + agree['first_intent_before_near']}/{len(cases)}")
    out.append(f"- 첫 송금 직전: 세 모델 완전 일치 {agree['first_transfer_before']}/{len(cases)}, ±1 포함 {agree['first_transfer_before'] + agree['first_transfer_before_near']}/{len(cases)}")

    with open(os.path.join(args.out, "report.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(out) + "\n")

    with open(os.path.join(args.out, "majority_messages.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["case", "seq", "text", "signals_majority", "signals_split", "stage_majority", "stage_votes"]
                   + [f"signals_{n}" for n in names] + ([f"signals_{human[0]}", f"stage_{human[0]}"] if human else []))
        for k in common:
            m = maj[k]
            row = [k[0], k[1], sheet_keys[k], ";".join(sorted(m["signals"])), ";".join(m["split_sig"]),
                   m["stage"], " ".join(f"{s}{v}" for s, v in m["stage_votes"].items())]
            row += [";".join(sorted(models[n][0][k]["signals"])) for n in names]
            if human:
                h = human[1].get(k)
                row += [";".join(sorted(h["signals"])) if h else "", h["stage"] if h else ""]
            w.writerow(row)

    with open(os.path.join(args.out, "disagreements.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["case", "seq", "text", "kind", "detail"] + [f"note_{n}" for n in names])
        for k in common:
            sets = {n: models[n][0][k]["signals"] for n in names}
            if len({frozenset(s) for s in sets.values()}) > 1:
                w.writerow([k[0], k[1], sheet_keys[k], "신호",
                            " / ".join(f"{n}:{';'.join(sorted(s)) or '-'}" for n, s in sets.items())]
                           + [models[n][0][k]["note"] for n in names])
            if not maj[k]["stage"] or len(maj[k]["stage_votes"]) > 1:
                w.writerow([k[0], k[1], sheet_keys[k], "단계",
                            " / ".join(f"{n}:{models[n][0][k]['stage']}" for n in names)]
                           + [models[n][0][k]["note"] for n in names])
    print(f"점검 문제 {len(problems)}개 → {os.path.join(args.out, 'report.md')}")


if __name__ == "__main__":
    main()
